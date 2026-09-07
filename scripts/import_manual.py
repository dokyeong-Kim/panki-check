#!/usr/bin/env python3
"""수동으로 내려받은 원자료를 표준 CSV로 바꾼다.

데이터랩 API가 막힌 환경에서는 사용자가 데이터랩 웹에서 xlsx를 내려받고,
기상자료개방포털에서 ASOS 일자료 CSV를 내려받아 올린다. 이 스크립트가 그
파일들을 찾아 data/search_daily.csv 와 data/weather_daily.csv 로 정규화한다.

사용법:
    python scripts/import_manual.py [파일 또는 폴더 ...]
인자가 없으면 /mnt/user-data/uploads 와 data/inbox 를 뒤진다.

- 데이터랩 xlsx: 시트 상단 6행 메타 + '날짜, 후드티, 날짜, 후드집업, 날짜, 플리스' 구조.
  기간을 넓혀 다시 받으면 정규화 기준이 바뀌므로 **항상 통째로 덮어쓴다**.
- ASOS CSV(cp949): 지점, 지점명, 일시, 평균기온, 최저기온, ..., 최고기온, ...
  여러 파일(연도별)을 한 번에 넣어도 된다. 기존 weather_daily.csv 와 병합하고
  같은 날짜는 새 파일이 이긴다.
"""
import csv
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
ITEMS = ["후드집업", "후드티", "플리스"]


def read_datalab_xlsx(path):
    try:
        from openpyxl import load_workbook
    except ImportError:
        sys.exit("openpyxl 이 필요합니다: pip install openpyxl")
    import warnings
    warnings.filterwarnings("ignore")
    # read_only 모드는 데이터랩 xlsx의 dimension 정보를 못 읽어 1행만 돌려준다
    wb = load_workbook(path, read_only=False, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    # 헤더 행 찾기: '날짜'로 시작하는 첫 행
    hi = next(i for i, r in enumerate(rows) if r and str(r[0]).strip() == "날짜")
    header = [str(c).strip() if c is not None else "" for c in rows[hi]]
    # 열 위치: 이름 → (날짜열, 값열)
    cols = {}
    for j, name in enumerate(header):
        if name in ITEMS:
            cols[name] = (j - 1, j)
    missing = [i for i in ITEMS if i not in cols]
    if missing:
        sys.exit(f"데이터랩 파일에 {missing} 열이 없습니다. 헤더: {header}")
    table = {}
    for r in rows[hi + 1:]:
        if not r or r[0] is None:
            continue
        for name, (dj, vj) in cols.items():
            d, v = r[dj], r[vj]
            if d is None or v is None:
                continue
            key = str(d)[:10]
            table.setdefault(key, {})[name] = float(v)
    return table


def read_asos_csv(path):
    out = {}
    for enc in ("cp949", "utf-8-sig", "utf-8"):
        try:
            with open(path, encoding=enc, newline="") as f:
                rd = csv.DictReader(f)
                for row in rd:
                    row = {k.strip(): (v or "").strip() for k, v in row.items() if k}
                    d = row.get("일시")
                    if not d:
                        continue
                    rec = {}
                    for src, dst in (("평균기온(°C)", "tavg"), ("최저기온(°C)", "tmin"), ("최고기온(°C)", "tmax")):
                        v = row.get(src, "")
                        if v not in ("", None):
                            rec[dst] = float(v)
                    if rec:
                        out[d] = rec
            return out
        except UnicodeDecodeError:
            out = {}
            continue
    sys.exit(f"인코딩을 못 읽음: {path}")


def load_weather_daily():
    p = DATA / "weather_daily.csv"
    out = {}
    if p.exists():
        with p.open(encoding="utf-8-sig", newline="") as f:
            for row in csv.DictReader(f):
                rec = {k: float(row[k]) for k in ("tmin", "tmax", "tavg") if row.get(k) not in ("", None)}
                out[row["date"]] = rec
    return out


def write_csv(path, table, keys):
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(["date"] + keys)
        for d in sorted(table):
            w.writerow([d] + [table[d].get(k, "") for k in keys])


def main(argv):
    targets = [Path(a) for a in argv] or [Path("/mnt/user-data/uploads"), DATA / "inbox"]
    files = []
    for t in targets:
        if t.is_dir():
            files += sorted(t.glob("*"))
        elif t.exists():
            files.append(t)
    xlsx = [f for f in files if f.suffix.lower() == ".xlsx"]
    asos = [f for f in files if f.suffix.lower() == ".csv" and f.name.upper().startswith("OBS_ASOS")]

    DATA.mkdir(exist_ok=True)
    if xlsx:
        latest = max(xlsx, key=lambda p: p.stat().st_mtime)
        table = read_datalab_xlsx(latest)
        write_csv(DATA / "search_daily.csv", table, ITEMS)
        ds = sorted(table)
        print(f"검색지수: {latest.name} → search_daily.csv  {ds[0]} ~ {ds[-1]} ({len(ds)}일)  [전체 덮어씀]")
    else:
        print("검색지수 xlsx 없음 — search_daily.csv 유지")

    if asos:
        wx = load_weather_daily()
        n0 = len(wx)
        for f in asos:
            wx.update(read_asos_csv(f))
        write_csv(DATA / "weather_daily.csv", wx, ["tmin", "tmax", "tavg"])
        ds = sorted(wx)
        print(f"기온: ASOS {len(asos)}개 파일 → weather_daily.csv  {ds[0]} ~ {ds[-1]} ({len(wx)}일, +{len(wx)-n0})")
    else:
        print("ASOS CSV 없음 — weather_daily.csv 유지")


if __name__ == "__main__":
    main(sys.argv[1:])
